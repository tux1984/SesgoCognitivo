import openpyxl
import pytest

from sesgocognitivo.corpus.excel_writer import (
    COLUMNA_AUX,
    COLUMNA_URL,
    FILA_INICIO_DATOS,
    encontrar_primera_fila_vacia,
    escribir_filas,
    guardar_workbook_seguro,
    leer_valores_permitidos,
)


class FilaFalsa:
    def __init__(self, **kw):
        self.__dict__.update(kw)


def _fila(**overrides):
    base = dict(
        articulo_id="A1",
        oracion_id="A1_S01",
        medio="Infobae Colombia",
        tema="Transición de gobierno y relación con la oposición",
        genero="Noticia dura",
        fecha_publicacion="2026-09-01",
        oracion_texto="Texto de prueba.",
        oracion_anterior="",
        oracion_siguiente="",
        url_fuente="https://example.com/nota",
    )
    base.update(overrides)
    return FilaFalsa(**base)


def test_primera_fila_vacia_es_8_no_7(fixture_wb_path):
    wb = openpyxl.load_workbook(fixture_wb_path)
    ws = wb["Corpus - oraciones"]
    assert encontrar_primera_fila_vacia(ws) == FILA_INICIO_DATOS == 8


def test_leer_valores_permitidos_lee_dropdown_real(fixture_wb_path):
    wb = openpyxl.load_workbook(fixture_wb_path)
    ws = wb["Corpus - oraciones"]
    assert "Infobae Colombia" in leer_valores_permitidos(ws, "C")
    assert "Transición de gobierno y relación con la oposición" in leer_valores_permitidos(ws, "D")
    assert set(leer_valores_permitidos(ws, "E")) == {"Noticia dura", "Opinión-análisis"}


def test_escribir_filas_empieza_en_8_y_no_toca_ejemplo(fixture_wb_path):
    wb = openpyxl.load_workbook(fixture_wb_path)
    ws = wb["Corpus - oraciones"]

    valores_antes = [[ws.cell(row=r, column=c).value for c in range(1, 33)] for r in (5, 6, 7)]

    escritas = escribir_filas(wb, [_fila()])
    assert escritas == 1
    assert ws["B8"].value == "A1_S01"
    assert ws["C8"].value == "Infobae Colombia"
    assert ws[f"{COLUMNA_URL}8"].value == "https://example.com/nota"

    valores_despues = [[ws.cell(row=r, column=c).value for c in range(1, 33)] for r in (5, 6, 7)]
    assert valores_antes == valores_despues


def test_escribir_filas_nunca_toca_columna_aux(fixture_wb_path):
    wb = openpyxl.load_workbook(fixture_wb_path)
    escribir_filas(wb, [_fila()])
    ws = wb["Corpus - oraciones"]
    valor = ws[f"{COLUMNA_AUX}8"].value
    assert isinstance(valor, str) and valor.startswith("=")


def test_escribir_filas_rechaza_medio_invalido(fixture_wb_path):
    wb = openpyxl.load_workbook(fixture_wb_path)
    with pytest.raises(ValueError):
        escribir_filas(wb, [_fila(medio="Medio inexistente")])


def test_configurar_columna_url_extiende_ambos_merges(fixture_wb_path):
    wb = openpyxl.load_workbook(fixture_wb_path)
    escribir_filas(wb, [_fila()])
    ws = wb["Corpus - oraciones"]
    rangos = {str(r) for r in ws.merged_cells.ranges}
    assert f"A1:{COLUMNA_URL}1" in rangos
    assert f"A2:{COLUMNA_URL}2" in rangos
    assert ws[f"{COLUMNA_URL}4"].value == "url_fuente"


def test_escribir_filas_es_idempotente_en_columna_url_setup(fixture_wb_path):
    """Llamar dos veces no debe duplicar el header ni romper el merge."""
    wb = openpyxl.load_workbook(fixture_wb_path)
    escribir_filas(wb, [_fila()])
    escribir_filas(wb, [_fila(oracion_id="A1_S02")])
    ws = wb["Corpus - oraciones"]
    assert ws[f"{COLUMNA_URL}4"].value == "url_fuente"
    assert ws["B9"].value == "A1_S02"


def test_escribir_filas_no_duplica_oracion_id_entre_corridas_separadas(fixture_wb_path):
    """Regresión de auditoría 2026-09-05: el corpus real terminó con 8 filas duplicadas
    (mismo oracion_id, mismo texto) porque el tracker en memoria evita sobre-contar una
    celda ya llena pero no evita GENERAR/ESCRIBIR una fila duplicada cuando el mismo
    artículo se reprocesa en una corrida SEPARADA (wb recién cargado) mientras la celda
    todavía tenía cupo. escribir_filas debe rechazar esto por sí solo, sin depender del
    tracker de la corrida que la llama."""
    wb = openpyxl.load_workbook(fixture_wb_path)
    escribir_filas(wb, [_fila(oracion_id="DUP_S01")])
    guardar_workbook_seguro(wb, fixture_wb_path)

    # Corrida "separada": se recarga el workbook desde disco, como haría una segunda
    # invocación real del CLI, y se intenta escribir el mismo oracion_id de nuevo.
    wb2 = openpyxl.load_workbook(fixture_wb_path)
    escritas = escribir_filas(wb2, [_fila(oracion_id="DUP_S01"), _fila(oracion_id="NUEVA_S01")])
    assert escritas == 1  # solo la fila nueva; la duplicada se descarta

    ws2 = wb2["Corpus - oraciones"]
    oracion_ids = [ws2.cell(row=r, column=2).value for r in range(8, 12)]
    assert oracion_ids.count("DUP_S01") == 1
    assert "NUEVA_S01" in oracion_ids


def test_escribir_filas_no_duplica_dentro_del_mismo_lote(fixture_wb_path):
    """Mismo chequeo, pero con el duplicado dentro de una única llamada (una corrida que
    por algún motivo genera el mismo oracion_id dos veces en su propia lista de filas)."""
    wb = openpyxl.load_workbook(fixture_wb_path)
    escritas = escribir_filas(wb, [_fila(oracion_id="X_S01"), _fila(oracion_id="X_S01")])
    assert escritas == 1


def test_guardar_workbook_seguro_ok(fixture_wb_path):
    wb = openpyxl.load_workbook(fixture_wb_path)
    escribir_filas(wb, [_fila()])
    guardar_workbook_seguro(wb, fixture_wb_path)

    releido = openpyxl.load_workbook(fixture_wb_path)
    assert releido["Corpus - oraciones"]["B8"].value == "A1_S01"
    assert not fixture_wb_path.with_suffix(".tmp.xlsx").exists()
