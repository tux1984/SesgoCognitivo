"""Escritura al Excel del corpus -- reglas duras, no solo convención:

  * Solo se escribe en columnas A-I (metadato/segmentación) y en la columna de url_fuente
    (ver COLUMNA_URL). NUNCA en J..AE ni en la columna de aux (ver COLUMNA_AUX): J..AE son
    de anotación humana (esquema completo de 9 categorías léxico/discursivo, ver
    docs/.../esquemas_anotacion.xlsx "Tabla 1"), la columna aux es la fórmula que alimenta
    el conteo de la hoja 'Grid de recoleccion'.
  * Las filas 5-7 de 'Corpus - oraciones' son un ejemplo protegido; los datos reales SIEMPRE
    empiezan en la fila 8 (fila_inicio=7 del script original era frágil: solo "funcionaba"
    porque la fila 7 no estaba vacía).
  * Cada valor de dropdown (medio/tema/género) se valida contra la lista real extraída de la
    propia validación de datos del archivo -- un drift entre `config/*.yaml` y el Excel se
    vuelve un ValueError explícito, no una fila corrupta silenciosa.
  * El guardado es atómico: se escribe a un .tmp.xlsx, se reabre y se hace un sanity check
    antes de reemplazar el archivo real.
"""
from __future__ import annotations

import logging
import os
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string

logger = logging.getLogger("sesgocognitivo")

HOJA_CORPUS = "Corpus - oraciones"
HOJAS_ESPERADAS = {"Grid de recoleccion", "Resumen", "Leyenda", "Corpus - oraciones"}

FILA_HEADER = 4
FILA_INICIO_DATOS = 8  # filas 5-7 = ejemplo protegido, nunca se tocan
FILA_FIN_DATOS = 557  # última fila pre-formateada / con validación de datos

COLUMNA_AUX = "AF"  # fórmula auxiliar de conteo (ver 'Grid de recoleccion'!H); NUNCA se edita
COLUMNA_URL = "AG"
COLUMNA_URL_HEADER = "url_fuente"

# Mapeo explícito campo -> columna. Deliberadamente NO usa enumerate/posición: así una
# columna nueva (url_fuente) no puede "correrse" hacia J (verbo_factivo, primera columna
# de anotación humana) por error.
COLUMNAS: dict[str, str] = {
    "articulo_id": "A",
    "oracion_id": "B",
    "medio": "C",
    "tema": "D",
    "genero": "E",
    "fecha_publicacion": "F",
    "oracion_texto": "G",
    "oracion_anterior": "H",
    "oracion_siguiente": "I",
    "url_fuente": COLUMNA_URL,
}


def _columna_de_sqref(sqref) -> str | None:
    primera = str(sqref).split(" ")[0].split(":")[0]
    letras = "".join(ch for ch in primera if ch.isalpha())
    return letras or None


def leer_valores_permitidos(ws, columna: str) -> list[str]:
    """Lee la lista real de un dropdown desde la propia validación de datos del archivo."""
    for dv in ws.data_validations.dataValidation:
        if dv.type == "list" and _columna_de_sqref(dv.sqref) == columna:
            formula = (dv.formula1 or "").strip('"')
            return [v for v in formula.split(",") if v]
    return []


def encontrar_primera_fila_vacia(ws, columna: str = "B", fila_inicio: int = FILA_INICIO_DATOS) -> int:
    fila_inicio = max(fila_inicio, FILA_INICIO_DATOS)
    col_idx = column_index_from_string(columna)
    fila = fila_inicio
    while ws.cell(row=fila, column=col_idx).value not in (None, ""):
        fila += 1
    return fila


def configurar_columna_url(wb) -> None:
    """Idempotente: agrega la columna de url_fuente (ver COLUMNA_URL) si todavía no existe.
    Copia el estilo de la columna aux (COLUMNA_AUX) para el header y extiende AMBOS títulos
    combinados hasta la nueva columna."""
    ws = wb[HOJA_CORPUS]
    header_cell = ws[f"{COLUMNA_URL}{FILA_HEADER}"]
    if header_cell.value == COLUMNA_URL_HEADER:
        return

    aux4 = ws[f"{COLUMNA_AUX}{FILA_HEADER}"]
    header_cell.value = COLUMNA_URL_HEADER
    header_cell.font = copy(aux4.font)
    header_cell.fill = copy(aux4.fill)
    header_cell.alignment = copy(aux4.alignment)
    header_cell.border = copy(aux4.border)

    rangos_actuales = {str(r) for r in ws.merged_cells.ranges}
    for fila_merge in (1, 2):
        rango_viejo = f"A{fila_merge}:{COLUMNA_AUX}{fila_merge}"
        rango_nuevo = f"A{fila_merge}:{COLUMNA_URL}{fila_merge}"
        if rango_viejo in rangos_actuales:
            ws.unmerge_cells(rango_viejo)
            ws.merge_cells(rango_nuevo)
            origen = ws[f"A{fila_merge}"]
            destino = ws[f"{COLUMNA_URL}{fila_merge}"]
            destino.font = copy(origen.font)
            destino.fill = copy(origen.fill)
            destino.alignment = copy(origen.alignment)

    ws.column_dimensions[COLUMNA_URL].width = 45
    logger.info("Columna %s ('%s') configurada en '%s'.", COLUMNA_URL, COLUMNA_URL_HEADER, HOJA_CORPUS)


def leer_oracion_ids_existentes(ws, fila_inicio: int = FILA_INICIO_DATOS, fila_fin: int = FILA_FIN_DATOS) -> set:
    """Todos los oracion_id ya presentes en la hoja (columna B), sin importar medio/tema/
    género. Es la fuente de verdad para deduplicar al escribir -- más fuerte que confiar
    solo en el tracker en memoria de la corrida actual (ver hallazgo de auditoría 2026-09-05:
    el tracker evita sobre-contar una celda ya llena, pero no evita GENERAR ni ESCRIBIR una
    fila duplicada cuando el mismo artículo se reprocesa en una corrida separada mientras su
    celda todavía tenía cupo -- así entraron 8 filas duplicadas de El Tiempo al corpus real,
    detectadas y corregidas en esa auditoría)."""
    col_idx = column_index_from_string("B")
    return {
        ws.cell(row=r, column=col_idx).value
        for r in range(fila_inicio, fila_fin + 1)
        if ws.cell(row=r, column=col_idx).value not in (None, "")
    }


def escribir_filas(wb, filas: list, fila_inicio_minima: int = FILA_INICIO_DATOS) -> int:
    """Escribe `filas` (objetos con atributos = claves de COLUMNAS) empezando en la primera
    fila libre >= fila_inicio_minima (nunca < 8). Nunca escribe en las columnas de anotación
    humana ni en la de aux (ver COLUMNA_AUX). Nunca escribe un
    oracion_id que ya exista en la hoja (ver `leer_oracion_ids_existentes`), sin importar de
    qué corrida/medio/tema venga -- red de seguridad dura contra duplicados, no solo una
    convención confiada al tracker en memoria. Devuelve cuántas filas se escribieron
    realmente (puede ser menos que len(filas) si se llega a la fila 557 o si algunas venían
    duplicadas)."""
    ws = wb[HOJA_CORPUS]
    configurar_columna_url(wb)

    valores_medio = set(leer_valores_permitidos(ws, "C"))
    valores_tema = set(leer_valores_permitidos(ws, "D"))
    valores_genero = set(leer_valores_permitidos(ws, "E"))
    oracion_ids_existentes = leer_oracion_ids_existentes(ws)

    fila = encontrar_primera_fila_vacia(ws, columna="B", fila_inicio=fila_inicio_minima)
    escritas = 0
    for f in filas:
        if f.oracion_id in oracion_ids_existentes:
            logger.warning(
                "Fila NO escrita: oracion_id ya existe en la hoja (duplicado): %s", f.oracion_id
            )
            continue
        if f.medio not in valores_medio:
            raise ValueError(f"medio inválido para el dropdown de la hoja: {f.medio!r}")
        if f.tema not in valores_tema:
            raise ValueError(f"tema inválido para el dropdown de la hoja: {f.tema!r}")
        if f.genero not in valores_genero:
            raise ValueError(f"genero inválido para el dropdown de la hoja: {f.genero!r}")
        if fila > FILA_FIN_DATOS:
            logger.error(
                "Se alcanzó la fila %d (límite pre-formateado del archivo). Deteniendo "
                "escritura; %d fila(s) de esta corrida NO se escribieron.",
                FILA_FIN_DATOS,
                len(filas) - escritas,
            )
            break

        estilo_texto = ws[f"G{fila}"]  # estilo de referencia (amarillo, wrap) para la fila nueva
        for campo, columna in COLUMNAS.items():
            valor = getattr(f, campo)
            celda = ws[f"{columna}{fila}"]
            celda.value = valor
            if columna == COLUMNA_URL:
                celda.font = copy(estilo_texto.font)
                celda.fill = copy(estilo_texto.fill)
                celda.alignment = copy(estilo_texto.alignment)
                celda.border = copy(estilo_texto.border)
        oracion_ids_existentes.add(f.oracion_id)  # por si `filas` trae duplicados internos
        fila += 1
        escritas += 1
    return escritas


def guardar_workbook_seguro(wb, ruta: Path) -> None:
    """Guardado atómico: escribe a un .tmp.xlsx, reabre y hace un sanity check, y solo
    entonces reemplaza el archivo original (os.replace, atómico a nivel de filesystem)."""
    ruta = Path(ruta)
    tmp = ruta.with_suffix(".tmp.xlsx")
    wb.save(tmp)

    verificacion = openpyxl.load_workbook(tmp, data_only=False)
    try:
        if set(verificacion.sheetnames) != HOJAS_ESPERADAS:
            raise RuntimeError(f"Sanity check falló: hojas inesperadas {verificacion.sheetnames}")
        ws_check = verificacion[HOJA_CORPUS]
        aux_muestra = ws_check[f"{COLUMNA_AUX}{FILA_INICIO_DATOS}"].value
        if not (isinstance(aux_muestra, str) and aux_muestra.startswith("=")):
            raise RuntimeError(f"Sanity check falló: la columna {COLUMNA_AUX} ya no contiene una fórmula.")
        # Filas 5-7 (ejemplo protegido) deben seguir teniendo su articulo_id original.
        if ws_check["A5"].value != "ART_0231" or ws_check["A7"].value != "ART_0245":
            raise RuntimeError("Sanity check falló: las filas de ejemplo (5-7) fueron modificadas.")
    except Exception:
        verificacion.close()
        tmp.unlink(missing_ok=True)
        raise
    verificacion.close()

    os.replace(tmp, ruta)
    logger.info("Workbook guardado de forma segura en %s", ruta)
